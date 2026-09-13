# -*- coding: utf-8 -*-
"""プラチナプログラム_自動化マスタ.xlsx を生成する"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date

WB = openpyxl.Workbook()

FONT = 'Arial'
H_FILL   = PatternFill('solid', fgColor='D9D9D9')   # 見出し
IN_FILL  = PatternFill('solid', fgColor='FFF2CC')   # 🟡 入力してください
AUTO_FILL= PatternFill('solid', fgColor='E2EFDA')   # 🟢 自動で入る
EX_FILL  = PatternFill('solid', fgColor='F2F2F2')   # 記入例
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def style_header(ws, row=1, ncols=None):
    ncols = ncols or ws.max_column
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, size=11)
        cell.fill = H_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)

def base_font(ws):
    for row in ws.iter_rows():
        for cell in row:
            if cell.font is None or cell.font.name != FONT:
                cell.font = Font(name=FONT, size=11,
                                 bold=bool(cell.font and cell.font.bold))

def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w

# ══════════════════════════════════════════════════════════
# ① 使い方
# ══════════════════════════════════════════════════════════
ws = WB.active
ws.title = '使い方'
rows = [
    ('', ''),
    ('プラチナプログラム 自動化マスタ', ''),
    ('', ''),
    ('このファイルの役割',
     'プラチナプログラムの運営を、Gmail下書きの自動作成＋活動報告まとめの自動作成＋\n'
     'Chatwork通知＋フォーム自動開閉 で回すためのマスタデータです。\n'
     'GAS（Google Apps Script）が毎朝このファイルを読んで動きます。'),
    ('', ''),
    ('シート構成', ''),
    ('① 基本設定', 'Zoom・フォームURL・Chatwork・署名など、毎回変わらない値の置き場。'),
    ('② スケジュール', 'プラチナグルコンの日程を1行1件で管理。GASが毎朝ここを読みます。'),
    ('③ メールテンプレート', 'タイミング（3日前／前日／当日／アーカイブ）ごとのメール雛形。\n'
                              '{{差込項目}} が自動で置き換わります。'),
    ('④ 受講生名簿', 'プラチナ受講生の氏名とメールアドレス。未予約者の判定に使います。'),
    ('⑤ 個別コンサル管理', '月1回の個別コンサルの予約案内・リマインドを管理します。'),
    ('⑥ 差込項目一覧', 'メールテンプレートで使える {{差込項目}} の一覧（参照用）。'),
    ('', ''),
    ('セルの色分け', ''),
    ('🟡 黄色', 'あなたが入力する場所'),
    ('🟢 緑', 'GASが自動で書き込む場所（触らない）'),
    ('⬜ グレー', '記入例。実際に使うときは削除するか、上書きしてください'),
    ('', ''),
    ('自動化の流れ（プラチナグルコン）', ''),
    ('5日前 朝', 'Gmailに「3日前」「前日」「当日」の下書きを3通作成 → Chatworkに通知\n'
                 '※宛先は入っていません。送信予約のときに手動で入れてください。'),
    ('前日 13:00', '事前フォームを自動でクローズ'),
    ('前日 13:30', '活動報告まとめドキュメントを作成 → Chatworkにリンクを通知\n'
                   '　同時に、スケジュールの「活動報告まとめURL」列に自動で書き込みます'),
    ('アーカイブ', 'スケジュールの「アーカイブ動画URL」を入力すると、翌朝にアーカイブメールの\n'
                   '下書きができます（メニューから手動でも作れます）'),
    ('', ''),
    ('最初にやること', ''),
    ('① 基本設定の🟡を埋める', 'Chatworkトークン、フォームURL、署名など'),
    ('② スケジュールに日程を入れる', '内容は必ず「プラチナグルコン」と入力してください（完全一致で判定します）'),
    ('③ 受講生名簿を入れる', '氏名とメールアドレス'),
    ('④ GASを設置する', 'このスプレッドシートの 拡張機能 → Apps Script にコードを貼り付け'),
]
for r, (a, b) in enumerate(rows, start=1):
    ws.cell(row=r, column=2, value=a)
    ws.cell(row=r, column=3, value=b)
ws['B2'].font = Font(name=FONT, bold=True, size=16)
for r in (4, 6, 14, 19, 26):
    ws.cell(row=r, column=2).font = Font(name=FONT, bold=True, size=12)
for r in range(1, len(rows) + 1):
    ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(row=r, column=2).alignment = Alignment(vertical='top')
widths(ws, {'A': 3, 'B': 30, 'C': 95})
base_font(ws)

print('使い方 OK')

from datetime import date
from templates import TEMPLATES

# ══════════════════════════════════════════════════════════
# ② 基本設定
# ══════════════════════════════════════════════════════════
ws = WB.create_sheet('基本設定')
ws.append(['項目', '設定値', 'メモ'])
SETTINGS = [
    ('プログラム名', 'プラチナプログラム', ''),
    ('', '', ''),
    ('── Zoom（毎回同じ）──', '', ''),
    ('プラチナグルコン ZoomURL', 'https://us02web.zoom.us/j/89349982545', '🟡'),
    ('プラチナグルコン ミーティングID', '893 4998 2545', '🟡'),
    ('', '', ''),
    ('── 事前フォーム ──', '', ''),
    ('プラチナグルコン事前フォームURL', 'https://ws.formzu.net/dist/S28714862/',
     '🟡 ★Googleフォームに変更したら、新しい回答用URLに差し替えてください'),
    ('', '', ''),
    ('── Chatwork設定 ──', '', ''),
    ('Chatwork APIトークン', '（GASのスクリプトプロパティに保存します。ここは空のままでOK）', ''),
    ('Chatwork ルームID', '444552021', '🟡 通知先ルーム'),
    ('Chatwork メンション先', '[To:5925962] 藤井みあさん', '🟡'),
    ('', '', ''),
    ('── 差出人・署名 ──', '', ''),
    ('事務局差出人名', 'Shine A Light 運営事務局', ''),
    ('事務局アドレス', 'shinealightonlineschool@gmail.com',
     '🟡 ★このアカウントでGASを作らないと、下書きがここに入りません'),
    ('署名', 'Shine A Light 運営事務局', '🟡 メール末尾の {{署名}} に入ります'),
    ('', '', ''),
    ('── メール本文で使う固定テキスト ──', '', ''),
    ('合宿予定',
     '＜合宿＞\n11月：東京　六本木付近\n12日（木）13日（金）',
     '🟡 アーカイブメールの {{合宿予定}} に入ります。変わったら書き換えてください'),
    ('', '', ''),
    ('── GAS実行時刻 ──', '', ''),
    ('下書き作成時刻', '07:00', '毎朝この時刻にGASが動きます'),
    ('下書きを作る日', '5', '🟡 グルコンの何日前に下書きを作るか（既定：5日前）'),
    ('活動報告フォームを開く日', '5', '🟡 グルコンの何日前にフォームを開くか'),
    ('フォーム締切時刻', '13:00', '🟡 前日のこの時刻にフォームを閉じます（案内は12時）'),
    ('まとめ送信時刻', '13:30', '🟡 前日のこの時刻にChatworkへ通知します'),
]
for row in SETTINGS:
    ws.append(list(row))
style_header(ws)
for r in range(2, ws.max_row + 1):
    a = ws.cell(row=r, column=1)
    if isinstance(a.value, str) and a.value.startswith('──'):
        a.font = Font(name=FONT, bold=True, size=11)
    memo = ws.cell(row=r, column=3).value or ''
    if '🟡' in str(memo):
        ws.cell(row=r, column=2).fill = IN_FILL
    for c in range(1, 4):
        ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical='top')
        ws.cell(row=r, column=c).border = BORDER
widths(ws, {'A': 34, 'B': 62, 'C': 62})
base_font(ws)

# ══════════════════════════════════════════════════════════
# ③ スケジュール
# ══════════════════════════════════════════════════════════
ws = WB.create_sheet('スケジュール')
HEAD = ['内容', '日程', '日程短', '開始時間', '終了時間',
        '下書き作成日', '提出期限', '質問まとめ',
        'アーカイブ動画URL', '活動報告まとめURL', 'ステータス']
ws.append(HEAD)
SCHEDULE = [
    (date(2026, 9, 1),  '10:00', '11:00', '実施済み'),
    (date(2026, 9, 15), '10:00', '11:00', '未実施'),
    (date(2026, 10, 6), '10:00', '11:00', '未実施'),
    (date(2026, 10, 27),'21:00', '22:00', '未実施'),
    (date(2026, 11, 10),'10:00', '11:00', '未実施'),
    (date(2026, 11, 24),'21:00', '22:00', '未実施'),
    (date(2026, 12, 1), '21:00', '22:00', '未実施'),
    (date(2026, 12, 15),'10:00', '11:00', '未実施'),
]
for i, (d, st, en, status) in enumerate(SCHEDULE):
    r = i + 2
    ws.cell(row=r, column=1, value='プラチナグルコン')
    ws.cell(row=r, column=2, value=d).number_format = 'yyyy"年"m"月"d"日"'
    ws.cell(row=r, column=3, value=f'=IF(B{r}="","",TEXT(B{r},"M/d（aaa）"))')
    ws.cell(row=r, column=4, value=st)
    ws.cell(row=r, column=5, value=en)
    ws.cell(row=r, column=6, value=f'=IF(B{r}="","",B{r}-$B$1)')   # 差し替え後に修正
    ws.cell(row=r, column=7, value=f'=IF(B{r}="","",B{r}-1)')
    ws.cell(row=r, column=8, value=f'=IF(B{r}="","",B{r}-1)')
    ws.cell(row=r, column=11, value=status)
    for c in (6, 7, 8):
        ws.cell(row=r, column=c).number_format = 'yyyy"年"m"月"d"日"'
    for c in (9, 10):
        ws.cell(row=r, column=c).fill = AUTO_FILL if c == 10 else IN_FILL
    for c in range(1, len(HEAD) + 1):
        ws.cell(row=r, column=c).border = BORDER
# 下書き作成日は「5日前」を素直に書く（基本設定を参照せず、見て分かる式に）
for i in range(len(SCHEDULE)):
    r = i + 2
    ws.cell(row=r, column=6, value=f'=IF(B{r}="","",B{r}-5)')
style_header(ws)
widths(ws, {'A': 20, 'B': 17, 'C': 14, 'D': 11, 'E': 11, 'F': 17,
            'G': 17, 'H': 17, 'I': 46, 'J': 46, 'K': 12})
base_font(ws)
note = ws.max_row + 2
ws.cell(row=note, column=1,
        value='※「内容」は必ず「プラチナグルコン」と入力（完全一致で判定します）')
ws.cell(row=note + 1, column=1,
        value='※🟡アーカイブ動画URL＝手入力　🟢活動報告まとめURL＝GASが自動で書き込みます')
ws.cell(row=note + 2, column=1,
        value='※C・F・G・H列は自動計算です（触らないでください）')
for r in (note, note + 1, note + 2):
    ws.cell(row=r, column=1).font = Font(name=FONT, size=10, italic=True)

print('基本設定・スケジュール OK')

# ══════════════════════════════════════════════════════════
# ④ メールテンプレート
# ══════════════════════════════════════════════════════════
ws = WB.create_sheet('メールテンプレート')
ws.append(['内容', 'タイミング', '件名', '本文'])
for row in TEMPLATES:
    ws.append(list(row))
style_header(ws)
for r in range(2, ws.max_row + 1):
    for c in range(1, 5):
        cell = ws.cell(row=r, column=c)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        cell.border = BORDER
    ws.cell(row=r, column=3).fill = IN_FILL
    ws.cell(row=r, column=4).fill = IN_FILL
    ws.row_dimensions[r].height = 120
widths(ws, {'A': 20, 'B': 13, 'C': 66, 'D': 100})
base_font(ws)

# ══════════════════════════════════════════════════════════
# ⑤ 受講生名簿
# ══════════════════════════════════════════════════════════
ws = WB.create_sheet('受講生名簿')
ws.append(['お名前', 'メールアドレス', '在籍状況', '備考'])
ws.append(['山田 花子', 'hanako.yamada@example.com', '在籍', '← 記入例。削除するか上書きしてください'])
style_header(ws)
for c in range(1, 5):
    ws.cell(row=2, column=c).fill = EX_FILL
    ws.cell(row=2, column=c).font = Font(name=FONT, size=11, italic=True, color='808080')
for r in range(3, 40):
    for c in range(1, 4):
        ws.cell(row=r, column=c).fill = IN_FILL
        ws.cell(row=r, column=c).border = BORDER
for c in range(1, 5):
    ws.cell(row=2, column=c).border = BORDER
widths(ws, {'A': 22, 'B': 34, 'C': 13, 'D': 48})
base_font(ws)
n = 41
ws.cell(row=n, column=1, value='※「在籍状況」が「在籍」の人だけが、メールとリマインドの対象になります')
ws.cell(row=n, column=1).font = Font(name=FONT, size=10, italic=True)

# ══════════════════════════════════════════════════════════
# ⑥ 個別コンサル管理
# ══════════════════════════════════════════════════════════
ws = WB.create_sheet('個別コンサル管理')
HEAD = ['対象月', '案内送付日', '予約URL', '予約可能日', '予約締切日',
        '案内下書き', 'リマインド前日', 'リマインド当日', 'リマインド翌日', 'ステータス']
ws.append(HEAD)
ws.append(['2026-09', date(2026, 9, 1),
           'https://tebanasu.net/wakanaonlinelesson/30minutesconsulting',
           '14、16、17、18、24、25、28日', date(2026, 9, 10),
           '未作成', '未作成', '未作成', '未作成', '← 記入例'])
style_header(ws)
for c in range(1, len(HEAD) + 1):
    ws.cell(row=2, column=c).fill = EX_FILL
    ws.cell(row=2, column=c).font = Font(name=FONT, size=11, italic=True, color='808080')
    ws.cell(row=2, column=c).border = BORDER
for c in (2, 5):
    ws.cell(row=2, column=c).number_format = 'yyyy/m/d'
for r in range(3, 18):
    for c in range(1, 6):
        ws.cell(row=r, column=c).fill = IN_FILL
        ws.cell(row=r, column=c).border = BORDER
    for c in range(6, 11):
        ws.cell(row=r, column=c).fill = AUTO_FILL
        ws.cell(row=r, column=c).border = BORDER
widths(ws, {'A': 12, 'B': 14, 'C': 52, 'D': 30, 'E': 14,
            'F': 13, 'G': 15, 'H': 15, 'I': 15, 'J': 16})
base_font(ws)
n = 19
for i, t in enumerate([
    '※ クライアントさんからの連絡を、A〜E列に転記するだけです（🟡）。',
    '※ F〜J列（🟢）はGASが自動で更新します。触らないでください。',
    '※「案内送付日」の朝に、予約案内メールの下書きができます。',
    '※ 予約締切日の前日・当日・翌日に、未予約の方だけへのリマインド下書きができます。',
    '※ 未予約の判定方法は tebanasu の確認後に決めます（現時点では未実装）。',
]):
    ws.cell(row=n + i, column=1, value=t).font = Font(name=FONT, size=10, italic=True)

print('メールテンプレート・受講生名簿・個別コンサル管理 OK')

# ══════════════════════════════════════════════════════════
# ⑦ 差込項目一覧
# ══════════════════════════════════════════════════════════
ws = WB.create_sheet('差込項目一覧')
ws.append(['差込項目', '入る内容（例：9/15グルコンの場合）', 'どこから取るか'])
MERGE = [
    ('{{日程}}',            '9月15日（火）',        'スケジュール「日程」'),
    ('{{日程短}}',          '9/15（火）',           'スケジュール「日程」'),
    ('{{時間帯}}',          '10時〜11時',           'スケジュール「開始時間」「終了時間」'),
    ('{{開始時}}',          '10時',                 'スケジュール「開始時間」'),
    ('{{提出期限}}',        '9月14日（月）',        'グルコン日の前日（自動計算）'),
    ('{{提出期限短}}',      '9/14(月)',             'グルコン日の前日（自動計算）'),
    ('{{zoomリンク}}',      'https://us02web.zoom.us/j/89349982545', '基本設定'),
    ('{{ミーティングID}}',  '893 4998 2545',        '基本設定'),
    ('{{事前フォームURL}}', 'https://ws.formzu.net/dist/S28714862/', '基本設定'),
    ('{{署名}}',            'Shine A Light 運営事務局', '基本設定'),
    ('{{合宿予定}}',        '＜合宿＞\n11月：東京　六本木付近\n12日（木）13日（金）', '基本設定'),
    ('{{アーカイブ動画URL}}','https://vimeo.com/xxxxxxx',
     '🟡 スケジュール「アーカイブ動画URL」（手入力）'),
    ('{{活動報告まとめURL}}','https://docs.google.com/document/d/xxxxx/edit',
     '🟢 スケジュール「活動報告まとめURL」（GASが自動で入れます）'),
    ('{{今後の日程}}',
     '9月：\n15日 （火）  午前10時〜11時\n\n10月：\n6日 （火）  午前10時〜11時\n27日 （火）  21時〜22時',
     '🟢 スケジュールから自動生成（今日以降の回を月ごとに並べます）'),
]
for row in MERGE:
    ws.append(list(row))
style_header(ws)
for r in range(2, ws.max_row + 1):
    for c in range(1, 4):
        ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical='top')
        ws.cell(row=r, column=c).border = BORDER
    ws.cell(row=r, column=1).font = Font(name=FONT, size=11, bold=True)
widths(ws, {'A': 26, 'B': 50, 'C': 56})
base_font(ws)
n = ws.max_row + 2
for i, t in enumerate([
    '※ メールテンプレートの件名・本文に上の {{...}} を書くと、送信内容に合わせて自動で置き換わります。',
    '※ 「午前」は開始時刻が12時より前のときだけ付きます（{{今後の日程}}）。',
]):
    ws.cell(row=n + i, column=1, value=t).font = Font(name=FONT, size=10, italic=True)

WB.save('プラチナプログラム_自動化マスタ.xlsx')
print('保存しました')
